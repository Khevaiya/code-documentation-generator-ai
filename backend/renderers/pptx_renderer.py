from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
import subprocess
from io import BytesIO
from core.models import DocumentModel

# Check if Node.js is available for diagram conversion
try:
    subprocess.run(['node', '--version'], capture_output=True, check=True)
    HAS_NODEJS = True
except (subprocess.CalledProcessError, FileNotFoundError):
    HAS_NODEJS = False


def _svg_to_png_nodejs(svg_content: str) -> bytes:
    """Convert SVG to PNG using Node.js and sharp library."""
    if not HAS_NODEJS:
        return None
    
    try:
        backend_dir = Path(__file__).parent.parent  # Go up from renderers/ to backend/
        
        # Create temporary files in backend directory
        svg_path = backend_dir / 'temp_pptx_diagram.svg'
        png_path = backend_dir / 'temp_pptx_diagram.png'
        js_path = backend_dir / 'temp_pptx_convert.js'
        
        # Write SVG content
        with open(svg_path, 'w', encoding='utf-8') as f:
            f.write(svg_content)
        
        # Node.js script to convert SVG to PNG
        node_script = """
const sharp = require('sharp');
const fs = require('fs');

const svgBuffer = fs.readFileSync('temp_pptx_diagram.svg');

sharp(svgBuffer)
  .resize({ width: 800, height: 600, fit: 'inside', withoutEnlargement: true })
  .png()
  .toFile('temp_pptx_diagram.png')
  .then(() => {
    console.log('SVG converted to PNG for PPTX');
  })
  .catch(err => {
    console.error('Error:', err);
    process.exit(1);
  });
"""
        
        # Write Node.js script
        with open(js_path, 'w') as f:
            f.write(node_script)
        
        # Run Node.js script
        result = subprocess.run(['node', 'temp_pptx_convert.js'], capture_output=True, text=True, timeout=30, cwd=str(backend_dir))
        
        if result.returncode == 0 and png_path.exists():
            # Read PNG data
            with open(png_path, 'rb') as f:
                png_data = f.read()
            
            # Cleanup
            svg_path.unlink(missing_ok=True)
            png_path.unlink(missing_ok=True)
            js_path.unlink(missing_ok=True)
            
            return png_data
        else:
            return None
            
    except Exception as e:
        return None
    finally:
        # Cleanup files if they exist
        try:
            for path in [svg_path, png_path, js_path]:
                if 'path' in locals() and path.exists():
                    path.unlink()
        except:
            pass


def render_pptx(doc: DocumentModel, output_path: str) -> str:
    """Render DocumentModel to PPTX using python-pptx."""
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    
    # Colors
    dark_blue = RGBColor(30, 58, 95)      # 1E3A5F
    medium_blue = RGBColor(46, 89, 132)   # 2E5984
    light_blue = RGBColor(170, 204, 238)  # AACCEE
    light_gray = RGBColor(241, 245, 249)  # F1F5F9
    dark_gray = RGBColor(51, 51, 51)      # 333333
    
    # Title Slide
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Background shape
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = dark_blue
    bg_shape.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.2), Inches(8.4), Inches(1.5)
    )
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = doc.project_name
    title_para.font.name = 'Arial'
    title_para.font.size = Pt(40)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(255, 255, 255)
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(2.7), Inches(8.4), Inches(0.6)
    )
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = "Codebase Analysis Report"
    subtitle_para.font.name = 'Arial'
    subtitle_para.font.size = Pt(20)
    subtitle_para.font.color.rgb = light_blue
    
    # Date
    date_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(3.5), Inches(8.4), Inches(0.5)
    )
    date_frame = date_box.text_frame
    date_para = date_frame.paragraphs[0]
    date_para.text = doc.generated_at.strftime('%B %d, %Y')
    date_para.font.name = 'Arial'
    date_para.font.size = Pt(14)
    date_para.font.color.rgb = RGBColor(136, 170, 204)
    
    # Project Overview Slide
    slide = prs.slides.add_slide(slide_layout)
    
    # Overview title
    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.3), Inches(8.8), Inches(0.7)
    )
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Project Overview"
    title_para.font.name = 'Arial'
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = dark_blue
    
    # Statistics
    sa = doc.static_analysis
    stats = [
        ("Language", sa.primary_language or "Unknown"),
        ("Architecture", sa.architecture_pattern or "Unknown"),
        ("Test Files", str(len(sa.test_files))),
        ("API Routes", str(len(sa.api_routes))),
        ("CI/CD", "Yes" if sa.has_ci else "No"),
        ("Docker", "Yes" if sa.has_docker else "No"),
    ]
    
    for i, (label, value) in enumerate(stats):
        col = i % 3
        row = i // 3
        x = Inches(0.6 + col * 3.1)
        y = Inches(1.3 + row * 1.6)
        
        # Background box
        bg_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(1.3)
        )
        bg_box.fill.solid()
        bg_box.fill.fore_color.rgb = light_gray
        bg_box.line.fill.background()
        
        # Label
        label_box = slide.shapes.add_textbox(
            x, y + Inches(0.15), Inches(2.8), Inches(0.4)
        )
        label_frame = label_box.text_frame
        label_para = label_frame.paragraphs[0]
        label_para.text = label
        label_para.alignment = PP_ALIGN.CENTER
        label_para.font.name = 'Arial'
        label_para.font.size = Pt(10)
        label_para.font.color.rgb = RGBColor(100, 116, 139)
        
        # Value
        value_box = slide.shapes.add_textbox(
            x, y + Inches(0.5), Inches(2.8), Inches(0.6)
        )
        value_frame = value_box.text_frame
        value_para = value_frame.paragraphs[0]
        value_para.text = value
        value_para.alignment = PP_ALIGN.CENTER
        value_para.font.name = 'Arial'
        value_para.font.size = Pt(20)
        value_para.font.bold = True
        value_para.font.color.rgb = dark_blue
    
    # Frameworks & Libraries Slide (if any)
    if sa.frameworks:
        slide = prs.slides.add_slide(slide_layout)
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(0.3), Inches(8.8), Inches(0.7)
        )
        title_frame = title_box.text_frame
        title_para = title_frame.paragraphs[0]
        title_para.text = "Frameworks & Libraries"
        title_para.font.name = 'Arial'
        title_para.font.size = Pt(28)
        title_para.font.bold = True
        title_para.font.color.rgb = dark_blue
        
        # Framework list
        content_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.2), Inches(8.8), Inches(4)
        )
        content_frame = content_box.text_frame
        
        for i, fw in enumerate(sa.frameworks[:10]):  # Limit to 10 frameworks
            if i > 0:
                para = content_frame.add_paragraph()
            else:
                para = content_frame.paragraphs[0]
            
            para.text = f"• {fw.name} ({fw.category}) — {fw.version or 'N/A'}"
            para.font.name = 'Arial'
            para.font.size = Pt(13)
            para.font.color.rgb = dark_gray
    
    # Lens Result Slides
    for lr in doc.lens_results:
        # Main lens slide
        slide = prs.slides.add_slide(slide_layout)
        
        # Header background
        header_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2)
        )
        header_bg.fill.solid()
        header_bg.fill.fore_color.rgb = dark_blue
        header_bg.line.fill.background()
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(0.2), Inches(8.8), Inches(0.8)
        )
        title_frame = title_box.text_frame
        title_para = title_frame.paragraphs[0]
        title_para.text = lr.title
        title_para.font.name = 'Arial'
        title_para.font.size = Pt(28)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(255, 255, 255)
        title_para.alignment = PP_ALIGN.CENTER
        
        # Summary
        summary_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.5), Inches(8.8), Inches(3.8)
        )
        summary_frame = summary_box.text_frame
        summary_frame.word_wrap = True
        summary_frame.auto_size = None
        summary_para = summary_frame.paragraphs[0]
        summary_para.text = lr.summary[:400] + ("..." if len(lr.summary) > 400 else "")
        summary_para.font.name = 'Arial'
        summary_para.font.size = Pt(12)
        summary_para.font.color.rgb = dark_gray
        
        # Diagrams slide (if any)
        if lr.diagrams:
            slide = prs.slides.add_slide(slide_layout)
            
            # Title
            title_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(0.3), Inches(8.8), Inches(0.6)
            )
            title_frame = title_box.text_frame
            title_para = title_frame.paragraphs[0]
            title_para.text = f"{lr.title} — Diagrams"
            title_para.font.name = 'Arial'
            title_para.font.size = Pt(22)
            title_para.font.bold = True
            title_para.font.color.rgb = dark_blue
            
            # Add diagrams as images or text
            y_pos = Inches(1.2)
            for i, diagram in enumerate(lr.diagrams[:3]):  # Limit to 3 diagrams per slide
                if diagram.svg and HAS_NODEJS:
                    # Try to convert and embed SVG as image
                    png_data = _svg_to_png_nodejs(diagram.svg)
                    if png_data:
                        png_stream = BytesIO(png_data)
                        # Add image to slide
                        slide.shapes.add_picture(png_stream, Inches(0.6), y_pos, width=Inches(8), height=Inches(1.5))
                        y_pos += Inches(1.8)
                        continue
                
                # Fallback: add as text
                content_box = slide.shapes.add_textbox(
                    Inches(0.6), y_pos, Inches(8.8), Inches(1.2)
                )
                content_frame = content_box.text_frame
                content_frame.word_wrap = True
                
                para = content_frame.paragraphs[0]
                para.text = f"• {diagram.title} ({diagram.diagram_type})"
                para.font.name = 'Arial'
                para.font.size = Pt(12)
                para.font.color.rgb = dark_gray
                
                if diagram.description:
                    desc_para = content_frame.add_paragraph()
                    desc_para.text = f"  {diagram.description[:100]}"
                    desc_para.font.name = 'Arial'
                    desc_para.font.size = Pt(10)
                    desc_para.font.color.rgb = RGBColor(100, 116, 139)
                
                y_pos += Inches(1.3)
        
        # Findings slide (if any)
        if lr.findings:
            slide = prs.slides.add_slide(slide_layout)
            
            # Title
            title_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(0.3), Inches(8.8), Inches(0.6)
            )
            title_frame = title_box.text_frame
            title_para = title_frame.paragraphs[0]
            title_para.text = f"{lr.title} — Findings"
            title_para.font.name = 'Arial'
            title_para.font.size = Pt(22)
            title_para.font.bold = True
            title_para.font.color.rgb = dark_blue
            
            # Findings list
            content_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(1.1), Inches(8.8), Inches(4.2)
            )
            content_frame = content_box.text_frame
            content_frame.word_wrap = True
            content_frame.auto_size = None
            
            for i, finding in enumerate(lr.findings[:6]):  # Limit to 6 findings
                if i > 0:
                    para = content_frame.add_paragraph()
                else:
                    para = content_frame.paragraphs[0]
                
                severity = (finding.severity or "info").upper()
                title_text = f"[{severity}] {finding.title[:80]}"
                desc_text = finding.description[:180]
                
                para.text = f"• {title_text}" + (f"\n  {desc_text}" if desc_text else "")
                para.font.name = 'Arial'
                para.font.size = Pt(10)
                para.font.color.rgb = dark_gray
        
        # Gaps slide (if any)
        if lr.gaps:
            slide = prs.slides.add_slide(slide_layout)
            
            # Title
            title_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(0.3), Inches(8.8), Inches(0.6)
            )
            title_frame = title_box.text_frame
            title_para = title_frame.paragraphs[0]
            title_para.text = f"{lr.title} — Gaps"
            title_para.font.name = 'Arial'
            title_para.font.size = Pt(22)
            title_para.font.bold = True
            title_para.font.color.rgb = RGBColor(204, 68, 0)  # Orange color
            
            # Gaps list
            content_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(1.1), Inches(8.8), Inches(4.2)
            )
            content_frame = content_box.text_frame
            content_frame.word_wrap = True
            content_frame.auto_size = None
            
            for i, gap in enumerate(lr.gaps[:8]):  # Limit to 8 gaps
                if i > 0:
                    para = content_frame.add_paragraph()
                else:
                    para = content_frame.paragraphs[0]
                
                para.text = f"• {gap.area}: {gap.description[:150]}"
                para.font.name = 'Arial'
                para.font.size = Pt(11)
                para.font.color.rgb = dark_gray
    
    # Thank You Slide
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = dark_blue
    bg_shape.line.fill.background()
    
    # Thank you text
    thanks_box = slide.shapes.add_textbox(
        Inches(1), Inches(1.5), Inches(8), Inches(1.5)
    )
    thanks_frame = thanks_box.text_frame
    thanks_para = thanks_frame.paragraphs[0]
    thanks_para.text = "Thank You"
    thanks_para.alignment = PP_ALIGN.CENTER
    thanks_para.font.name = 'Arial'
    thanks_para.font.size = Pt(44)
    thanks_para.font.bold = True
    thanks_para.font.color.rgb = RGBColor(255, 255, 255)
    
    # Footer text
    footer_box = slide.shapes.add_textbox(
        Inches(1), Inches(3.2), Inches(8), Inches(0.5)
    )
    footer_frame = footer_box.text_frame
    footer_para = footer_frame.paragraphs[0]
    footer_para.text = "Report generated by CodeLens AI - Codebase Analysis Agent"
    footer_para.alignment = PP_ALIGN.CENTER
    footer_para.font.name = 'Arial'
    footer_para.font.size = Pt(14)
    footer_para.font.color.rgb = RGBColor(136, 170, 204)
    
    # Save presentation
    prs.save(output_path)
    return output_path
