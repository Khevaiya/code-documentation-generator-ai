from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from io import BytesIO
import subprocess
import tempfile
import os
from core.models import DocumentModel

# Check if Node.js is available
try:
    subprocess.run(['node', '--version'], capture_output=True, check=True)
    HAS_NODEJS = True
except (subprocess.CalledProcessError, FileNotFoundError):
    HAS_NODEJS = False


def _svg_to_png_nodejs(svg_content: str) -> bytes:
    """Convert SVG to PNG using Node.js and sharp library."""
    if not HAS_NODEJS:
        return None
    
    try:
        backend_dir = Path(__file__).parent.parent  # Go up from renderers/ to backend/
        
        # Create temporary files in backend directory
        svg_path = backend_dir / 'temp_diagram.svg'
        png_path = backend_dir / 'temp_diagram.png'
        js_path = backend_dir / 'temp_convert.js'
        
        # Write SVG content
        with open(svg_path, 'w', encoding='utf-8') as f:
            f.write(svg_content)
        
        # Node.js script to convert SVG to PNG
        node_script = f"""
const sharp = require('sharp');
const fs = require('fs');

const svgBuffer = fs.readFileSync('temp_diagram.svg');

sharp(svgBuffer)
  .resize({{ width: 600, height: 400, fit: 'inside', withoutEnlargement: true }})
  .png()
  .toFile('temp_diagram.png')
  .then(() => {{
    console.log('SVG converted to PNG successfully');
  }})
  .catch(err => {{
    console.error('Error:', err);
    process.exit(1);
  }});
"""
        
        # Write Node.js script
        with open(js_path, 'w') as f:
            f.write(node_script)
        
        # Run Node.js script from backend directory
        result = subprocess.run(['node', 'temp_convert.js'], capture_output=True, text=True, timeout=30, cwd=str(backend_dir))
        
        if result.returncode == 0 and png_path.exists():
            # Read PNG data
            with open(png_path, 'rb') as f:
                png_data = f.read()
            
            # Cleanup
            svg_path.unlink(missing_ok=True)
            png_path.unlink(missing_ok=True)
            js_path.unlink(missing_ok=True)
            
            return png_data
        else:
            print(f"[docx_renderer] Node.js conversion failed: {result.stderr}")
            return None
            
    except Exception as e:
        print(f"[docx_renderer] SVG conversion error: {e}")
        return None
    finally:
        # Cleanup files if they exist
        try:
            for path in [svg_path, png_path, js_path]:
                if 'path' in locals() and path.exists():
                    path.unlink()
        except:
            pass


def render_docx(doc: DocumentModel, output_path: str) -> str:
    """Render DocumentModel to DOCX using python-docx."""
    # Create document
    docx_doc = Document()
    
    # Set up styles
    styles = docx_doc.styles
    
    # Title style
    title_style = styles['Title']
    title_font = title_style.font
    title_font.name = 'Arial'
    title_font.size = Pt(24)
    title_font.color.rgb = RGBColor(30, 58, 95)  # Dark blue
    
    # Heading styles
    heading1_style = styles['Heading 1']
    heading1_font = heading1_style.font
    heading1_font.name = 'Arial'
    heading1_font.size = Pt(18)
    heading1_font.color.rgb = RGBColor(30, 58, 95)
    
    heading2_style = styles['Heading 2']
    heading2_font = heading2_style.font
    heading2_font.name = 'Arial'
    heading2_font.size = Pt(14)
    heading2_font.color.rgb = RGBColor(46, 89, 132)
    
    # Document title
    title = docx_doc.add_heading(doc.project_name, 0)
    title.style = title_style
    
    # Generated date
    date_para = docx_doc.add_paragraph(f'Generated: {doc.generated_at.strftime("%B %d, %Y")}')
    date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Project Overview
    docx_doc.add_heading('Project Overview', level=1)
    
    sa = doc.static_analysis
    overview_text = (
        f"Language: {sa.primary_language or 'Unknown'} | "
        f"Architecture: {sa.architecture_pattern or 'Unknown'} | "
        f"Tests: {'Yes' if sa.has_tests else 'No'} ({len(sa.test_files)} files) | "
        f"Docker: {'Yes' if sa.has_docker else 'No'} | "
        f"CI: {'Yes' if sa.has_ci else 'No'} | "
        f"API Routes: {len(sa.api_routes)}"
    )
    docx_doc.add_paragraph(overview_text)
    
    # Frameworks & Libraries
    if sa.frameworks:
        docx_doc.add_heading('Frameworks & Libraries', level=2)
        table = docx_doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.style = 'Table Grid'
        
        # Header row
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Framework'
        hdr_cells[1].text = 'Category'
        hdr_cells[2].text = 'Version'
        
        # Style header
        for cell in hdr_cells:
            cell.paragraphs[0].runs[0].font.bold = True
        
        # Data rows
        for fw in sa.frameworks:
            row_cells = table.add_row().cells
            row_cells[0].text = fw.name
            row_cells[1].text = fw.category
            row_cells[2].text = fw.version or 'N/A'
    
    # Page break
    docx_doc.add_page_break()
    
    # Lens Results
    for lr in doc.lens_results:
        docx_doc.add_heading(lr.title, level=1)
        docx_doc.add_paragraph(lr.summary)
        
        # Diagrams
        if lr.diagrams:
            docx_doc.add_heading('Diagrams', level=2)
            for diagram in lr.diagrams:
                diagram_para = docx_doc.add_paragraph(f'{diagram.title} ({diagram.diagram_type})')
                diagram_para.runs[0].font.bold = True
                diagram_para.runs[0].font.color.rgb = RGBColor(46, 89, 132)
                
                if diagram.svg and HAS_NODEJS:
                    try:
                        # Convert SVG to PNG using Node.js
                        png_data = _svg_to_png_nodejs(diagram.svg)
                        if png_data:
                            png_stream = BytesIO(png_data)
                            docx_doc.add_picture(png_stream, width=Inches(6))
                        else:
                            docx_doc.add_paragraph('[Diagram conversion failed - Node.js/sharp required]')
                    except Exception as e:
                        docx_doc.add_paragraph(f'[Diagram error: {str(e)}]')
                elif diagram.svg:
                    # Fallback when Node.js not available
                    if diagram.description:
                        docx_doc.add_paragraph(f'Description: {diagram.description}')
                    docx_doc.add_paragraph('[Visual diagram available in HTML and PDF formats - install Node.js and sharp for DOCX diagrams]')
                else:
                    docx_doc.add_paragraph('[Diagram unavailable]')
        
        # Findings
        if lr.findings:
            docx_doc.add_heading('Findings', level=2)
            for finding in lr.findings:
                # Severity colors
                severity_colors = {
                    'CRITICAL': RGBColor(204, 0, 0),
                    'HIGH': RGBColor(255, 102, 0),
                    'MEDIUM': RGBColor(255, 170, 0),
                    'LOW': RGBColor(0, 102, 204),
                }
                
                severity = (finding.severity or 'INFO').upper()
                color = severity_colors.get(severity, RGBColor(102, 102, 102))
                
                # Finding title with severity
                finding_para = docx_doc.add_paragraph()
                title_run = finding_para.add_run(f'[{severity}] {finding.title}')
                title_run.font.bold = True
                title_run.font.color.rgb = color
                
                # Finding description
                docx_doc.add_paragraph(finding.description)
                
                # Recommendation
                if finding.recommendation:
                    rec_para = docx_doc.add_paragraph(f'Recommendation: {finding.recommendation}')
                    rec_para.runs[0].font.italic = True
        
        # Gaps
        if lr.gaps:
            docx_doc.add_heading('Gaps & Missing Items', level=2)
            for gap in lr.gaps:
                gap_para = docx_doc.add_paragraph(f'{gap.area}: {gap.description}', style='List Bullet')
        
        # Raw sections
        if lr.raw_sections:
            docx_doc.add_heading('Additional Details', level=2)
            for section_name, section_content in lr.raw_sections.items():
                docx_doc.add_heading(section_name, level=3)
                
                # Flatten content
                content_lines = _flatten_content(section_content)
                for line in content_lines[:50]:  # Limit to prevent excessive length
                    if line.strip():
                        docx_doc.add_paragraph(line)
    
    # Footer
    footer_para = docx_doc.add_paragraph('\nGenerated by CodeLens AI - Codebase Analysis Agent')
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.runs[0].font.italic = True
    footer_para.runs[0].font.color.rgb = RGBColor(153, 153, 153)
    
    # Save document
    docx_doc.save(output_path)
    return output_path


def _flatten_content(content, depth=0, max_depth=3):
    """Flatten nested content into readable lines."""
    lines = []
    if depth > max_depth:
        return lines
        
    indent = '  ' * depth
    
    if isinstance(content, str):
        for line in content.split('\n'):
            if line.strip():
                lines.append(indent + line.strip())
    elif isinstance(content, list):
        for item in content:
            if isinstance(item, str):
                lines.append(indent + '• ' + item)
            else:
                lines.extend(_flatten_content(item, depth + 1, max_depth))
    elif isinstance(content, dict):
        for key, value in content.items():
            key_formatted = key.replace('_', ' ').title()
            if isinstance(value, (str, int, float)):
                lines.append(f'{indent}{key_formatted}: {value}')
            else:
                lines.append(f'{indent}{key_formatted}:')
                lines.extend(_flatten_content(value, depth + 1, max_depth))
    else:
        lines.append(indent + str(content))
    
    return lines
